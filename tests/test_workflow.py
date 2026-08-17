from __future__ import annotations

import re
from pathlib import Path

import yaml
from conftest import (
    FakeLiteratureService,
    GuardRetryProvider,
    ScriptedProvider,
    UAVRegressionProvider,
)
from docx import Document
from docx.shared import RGBColor

from paperforge.config import load_config
from paperforge.domain import (
    AuthorValidationDecision,
    EvidenceCoverageStatus,
    PaperType,
    ResearchProfile,
    StageStatus,
)
from paperforge.llm import LLMClient
from paperforge.stages import StageRunner
from paperforge.storage import ProjectStore
from paperforge.workflow import WorkflowEngine


def engine_for(store: ProjectStore, provider):
    config = load_config(store.config_path)
    literature = FakeLiteratureService()
    runner = StageRunner(store, config, LLMClient(config, provider), literature)
    return WorkflowEngine(store, config, runner), literature


def test_topic_only_input_completes_without_questions(
    project_store: ProjectStore,
) -> None:
    provider = ScriptedProvider()
    engine, literature = engine_for(project_store, provider)
    report = engine.run()
    state = project_store.load_state()
    assert report.records[-1].stage == "export"
    assert state.profile.resolved_paper_type == PaperType.REVIEW_ARTICLE
    assert state.workflow_completed is True
    assert state.submission_ready is False
    assert state.author_actions
    assert all(state.status_for(stage).complete for stage in engine.config.workflow.stages)
    assert not (project_store.root / "inputs" / "required_facts.yaml").exists()
    assert (project_store.root / "outputs" / "manuscript.md").exists()
    assert (project_store.root / "outputs" / "literature-matrix.csv").exists()
    assert (project_store.root / "outputs" / "publication-profile.md").exists()
    assert (project_store.root / "outputs" / "submission-checklist.md").exists()
    document = Document(project_store.root / "outputs" / "manuscript.docx")
    references = [
        paragraph.text
        for paragraph in document.paragraphs
        if re.match(r"^\d+\.\s+", paragraph.text)
    ]
    assert references[0].startswith("1. ")
    assert any(paragraph.text.startswith("1. Introduction") for paragraph in document.paragraphs)
    assert document.styles["Heading 2"].font.name == "Times New Roman"
    assert document.styles["Heading 2"].font.color.rgb == RGBColor(0, 0, 0)
    assert literature.calls == 1


def test_topic_only_with_model_number_is_not_misclassified(
    tmp_path: Path,
    default_config_path: Path,
) -> None:
    store = ProjectStore(tmp_path / "topic-with-number")
    store.initialize(
        ResearchProfile(topic="YOLOv8 method accuracy for Industry 4.0 thermal monitoring"),
        default_config_path,
    )
    payload = yaml.safe_load(store.config_path.read_text(encoding="utf-8"))
    payload["quality"].update(
        {
            "minimum_verified_sources": 3,
            "minimum_cited_sources": 3,
            "minimum_manuscript_words": 300,
            "minimum_tables_for_original_research": 0,
            "require_section_depth": False,
        }
    )
    payload["literature"].update({"min_sources": 3, "target_sources": 6, "max_sources": 8})
    payload["workflow"]["max_review_cycles"] = 1
    store.config_path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    engine, _ = engine_for(store, ScriptedProvider())

    engine.run()

    assert store.load_state().profile.resolved_paper_type == PaperType.REVIEW_ARTICLE


def test_completed_workflow_is_idempotent_and_input_change_rebuilds(
    project_store: ProjectStore,
) -> None:
    provider = ScriptedProvider()
    engine, _ = engine_for(project_store, provider)
    engine.run()
    calls_after_first = provider.calls.copy()
    second = engine.run()
    assert second.records == []
    assert provider.calls == calls_after_first
    assert second.resumed_stages == len(engine.config.workflow.stages)

    brief = project_store.root / "inputs" / "research_brief.md"
    brief.write_text(
        brief.read_text(encoding="utf-8") + "\nAuthentic scope clarification.\n",
        encoding="utf-8",
    )
    third = engine.run()
    assert third.invalidated is True
    assert provider.calls["plan"] == calls_after_first["plan"] + 1
    assert project_store.load_state().workflow_completed is True


def test_rejected_revision_is_preserved_and_retried(project_store: ProjectStore) -> None:
    provider = GuardRetryProvider()
    engine, _ = engine_for(project_store, provider)

    engine.run()

    record = project_store.load_state().stage_records["evidence_review"]
    assert record.status == StageStatus.PASSED
    assert provider.calls["revise"] >= 2
    assert any("Rejected 1 unsafe targeted revision" in change for change in record.changes)
    assert "deliberately incomplete revision" not in project_store.read_manuscript()


def test_incomplete_uav_study_researches_drafts_and_defers_one_validation_pass(
    tmp_path: Path,
    default_config_path: Path,
) -> None:
    store = ProjectStore(tmp_path / "uav-paper")
    store.initialize(
        ResearchProfile(
            topic="Edge-cloud UAV thermal fire detection and geo-tagged alerting",
            requested_paper_type=PaperType.AUTO,
            target_journal="DJES",
        ),
        default_config_path,
    )
    config_payload = yaml.safe_load(store.config_path.read_text(encoding="utf-8"))
    config_payload["quality"]["minimum_verified_sources"] = 3
    config_payload["quality"]["minimum_cited_sources"] = 3
    config_payload["quality"]["minimum_manuscript_words"] = 300
    config_payload["quality"]["require_section_depth"] = False
    config_payload["literature"]["min_sources"] = 3
    config_payload["literature"]["target_sources"] = 6
    config_payload["literature"]["max_sources"] = 8
    config_payload["workflow"]["max_review_cycles"] = 1
    store.config_path.write_text(
        yaml.safe_dump(config_payload, sort_keys=False),
        encoding="utf-8",
    )
    responses = {
        "answers": {
            "Q-001": (
                "The UAV system uses a Pixhawk Cube flight controller, Raspberry Pi 4 edge "
                "processor, Waveshare 80×62 LWIR thermal camera, and Here3+ GNSS module. Thermal "
                "fire detection is performed onboard using temperature thresholding, hotspot-area "
                "filtering, thermal-contrast analysis, connected-region evaluation, and temporal "
                "confirmation. Only confirmed fire events with timestamp, GNSS data, and thermal "
                "evidence are transmitted to the cloud over Wi-Fi."
            ),
            "Q-002": (
                "A labelled dataset of 1,000 thermal frames was used, comprising 500 fire and 500 "
                "non-fire images collected during controlled low-altitude experiments. Each frame "
                "was labelled as fire or non-fire. The current dataset represents controlled "
                "experimental locations rather than broad geographic coverage."
            ),
            "Q-003": (
                "The system achieved 96.80% accuracy, 96.43% precision, 97.20% recall, 96.40% "
                "specificity, 96.81% F1-score, 3.60% false-positive rate, 2.80% false-negative "
                "rate, and MCC = 0.936. Response-time testing showed accuracy decreasing from "
                "97.4% at 220 ms to 93.1% at 920 ms. Formal comparison with external baseline "
                "methods is not yet included."
            ),
            "Q-004": (
                "The system was validated through controlled hardware-based UAV experiments, not "
                "simulation. Thermal fire and non-fire frames were processed using the onboard "
                "detection pipeline, and predictions were compared with ground-truth labels. "
                "Results demonstrated high detection reliability with 96.8% overall accuracy and "
                "effective geo-tagged event generation."
            ),
            "Q-005": (
                "Deployment considerations include UAV battery endurance, payload weight, "
                "Raspberry Pi processing capability, thermal sensor resolution, communication "
                "coverage and latency, GNSS accuracy, operating altitude, environmental thermal "
                "variation, and flight regulations. The event-driven edge-cloud design reduces "
                "bandwidth and dependence on continuous connectivity, while temporary communication "
                "failures are handled through local buffering and retransmission."
            ),
        }
    }
    store.write_text(
        "inputs/responses.yaml",
        yaml.safe_dump(responses, sort_keys=False),
    )
    provider = UAVRegressionProvider()
    engine, literature = engine_for(store, provider)
    report = engine.run()
    state = store.load_state()
    evidence_mapping = state.stage_records["evidence_mapping"]

    assert report.records[-1].stage == "export"
    assert state.profile.resolved_paper_type == PaperType.ORIGINAL_RESEARCH
    assert evidence_mapping.status == StageStatus.PASSED_WITH_ACTIONS
    assert state.stage_records["author_validation"].status == StageStatus.PASSED_WITH_ACTIONS
    assert state.workflow_completed is True
    assert state.submission_ready is False
    codes = {issue.code for issue in evidence_mapping.issues}
    assert "unverified_study_detail_algorithm_parameters" in codes
    assert "unverified_study_detail_evaluation_independence" in codes
    assert "unverified_study_detail_statistical_support" in codes
    assert "unverified_study_detail_calibration" in codes
    assert store.read_manuscript().strip()
    assert (store.root / "outputs" / "manuscript.docx").exists()
    validation = store.load_author_validation()
    validation_codes = {item.requirement_code for item in validation.pending_items}
    assert "acquisition_protocol" in validation_codes
    assert "ground_truth_protocol" in validation_codes
    assert "algorithm_parameters" in validation_codes
    assert all(item.literature_context for item in validation.items)
    assert literature.calls == 1
    assert provider.calls["plan"] == 1
    assert provider.calls["synthesis"] == 1
    manuscript = store.read_manuscript().casefold()
    assert "12 flights" not in manuscript
    assert "two independent experts" not in manuscript
    calls_after_first = provider.calls.copy()
    second = engine.run()
    assert second.records == []
    assert second.invalidated is False
    assert provider.calls == calls_after_first

    acquisition = next(
        item for item in validation.items if item.requirement_code == "acquisition_protocol"
    )
    acquisition.decision = AuthorValidationDecision.PROVIDED
    acquisition.answer = (
        "The experiment comprised 6 flights at Site A at an altitude of 20 m; all frames were "
        "grouped by flight before evaluation."
    )
    store.save_author_validation(validation)
    third = engine.run()
    assert third.invalidated is True
    updated_coverage = store.load_evidence_coverage()
    updated_acquisition = next(
        item for item in updated_coverage.requirements if item.code == "acquisition_protocol"
    )
    assert updated_acquisition.status == EvidenceCoverageStatus.SUPPORTED
    assert any(
        evidence.id == "EV-AUTHOR-VALIDATION" and "6 flights at Site A" in evidence.content
        for evidence in store.load_evidence()
    )
    rebuilt_validation = store.load_author_validation()
    assert all(
        "search timestamp" not in fact.casefold() and "discovered records" not in fact.casefold()
        for item in rebuilt_validation.items
        for fact in item.known_facts
    )


def test_strict_pre_draft_mode_remains_available(
    tmp_path: Path,
    default_config_path: Path,
) -> None:
    store = ProjectStore(tmp_path / "strict-paper")
    store.initialize(
        ResearchProfile(
            topic="UAV thermal fire detection experiment",
            requested_paper_type=PaperType.ORIGINAL_RESEARCH,
        ),
        default_config_path,
    )
    store.write_text(
        "inputs/responses.yaml",
        yaml.safe_dump(
            {
                "answers": {
                    "Q-001": (
                        "A Raspberry Pi thermal-camera experiment evaluated 1,000 labelled "
                        "frames and achieved 96.8% accuracy."
                    )
                }
            },
            sort_keys=False,
        ),
    )
    payload = yaml.safe_load(store.config_path.read_text(encoding="utf-8"))
    payload["workflow"]["evidence_gap_mode"] = "strict_pre_draft"
    store.config_path.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    engine, literature = engine_for(store, ScriptedProvider())

    report = engine.run()

    assert report.records[-1].stage == "evidence_mapping"
    assert report.records[-1].status == StageStatus.BLOCKED
    assert literature.calls == 0
    assert not store.read_manuscript().strip()
