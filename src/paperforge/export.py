from __future__ import annotations

import csv
import io
import json
import re
from datetime import UTC, datetime
from html import escape

from paperforge.schemas import Section, Source
from paperforge.store import Store


def ieee(source: Source, number: int) -> str:
    authors = ", ".join(source.authors) if source.authors else "[AUTHOR METADATA REQUIRED]"
    parts = [f"[{number}] {authors}, “{source.title},”"]
    if source.venue:
        parts.append(source.venue + ",")
    if source.volume:
        parts.append(f"vol. {source.volume},")
    if source.issue:
        parts.append(f"no. {source.issue},")
    if source.pages:
        parts.append(f"pp. {source.pages},")
    parts.append(f"{source.year}." if source.year else "[YEAR REQUIRED].")
    parts.append(f"doi: {source.doi}." if source.doi else f"[Online]. Available: {source.url}")
    return " ".join(parts)


def assemble(topic: str, sections: list[Section], sources: list[Source]) -> tuple[str, list[str]]:
    numbers: dict[str, int] = {}
    registry = {source.id: source for source in sources}
    ordered = sorted(
        sections, key=lambda section: {"Abstract": -2, "Keywords": -1}.get(section.title, 0)
    )
    parts = ["# " + topic]
    for section in ordered:
        parts.append("## " + section.title)
        for paragraph in section.paragraphs:
            citations = []
            for source_id in dict.fromkeys(paragraph.source_ids):
                if source_id not in registry:
                    raise ValueError("Cannot export unknown citations")
                numbers.setdefault(source_id, len(numbers) + 1)
                citations.append(f"[{numbers[source_id]}]")
            parts.append(paragraph.text.strip() + (" " + ", ".join(citations) if citations else ""))
    if numbers:
        parts.extend(
            ["## References"]
            + [ieee(registry[identifier], number) for identifier, number in numbers.items()]
        )
    return "\n\n".join(parts) + "\n", list(numbers)


def svg_diagram(nodes: list[str], edges: list[tuple[int, int]], title: str) -> str:
    height = 100 + len(nodes) * 90
    body = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="760" height="{height}" viewBox="0 0 760 {height}">',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#334155"/></marker></defs>',
        f'<rect width="100%" height="100%" fill="white"/><text x="380" y="30" text-anchor="middle" font-family="Arial" font-size="22">{escape(title)}</text>',
    ]
    for start, end in edges:
        if start < 0 or end < 0 or start >= len(nodes) or end >= len(nodes) or start == end:
            raise ValueError("Diagram edge references an invalid node")
        if end == start + 1:
            body.append(
                f'<path d="M380,{75 + start * 90 + 50} L380,{75 + end * 90}" stroke="#334155" fill="none" stroke-width="2" marker-end="url(#arrow)"/>'
            )
        else:
            side = 70 if end > start else 690
            body.append(
                f'<path d="M580,{100 + start * 90} H{side} V{100 + end * 90} H580" stroke="#334155" fill="none" stroke-width="2" marker-end="url(#arrow)"/>'
            )
    for index, label in enumerate(nodes):
        body.append(
            f'<rect x="180" y="{75 + index * 90}" width="400" height="50" rx="6" fill="#f1f5f9" stroke="#334155"/>'
        )
        # Keep labels readable; full text remains in the separate figure specification.
        lines = [label[i : i + 42] for i in range(0, min(len(label), 84), 42)]
        for row, line in enumerate(lines):
            body.append(
                f'<text x="380" y="{95 + index * 90 + row * 20}" text-anchor="middle" font-family="Arial" font-size="18">{escape(line)}</text>'
            )
    body.append("</svg>")
    return "".join(body)


def export_packet(
    store: Store,
    sections: list[Section],
    sources: list[Source],
    results: list[dict],
    actions: list[str],
    *,
    review: dict | None,
    recent_years: int,
    plan: dict | None = None,
    assessments: list[dict] | None = None,
) -> dict:
    project = store.get("project")
    manuscript, cited = assemble(
        (plan or {}).get("manuscript_title") or project["topic"], sections, sources
    )
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%S%f")
    files = [
        store.write("outputs/manuscript.md", manuscript),
        store.write(f"audit/manuscript-{stamp}.md", manuscript),
    ]
    full_actions = list(dict.fromkeys(actions))
    if review:
        full_actions.extend(
            finding["required_change"] for finding in review["issues"] if finding["blocking"]
        )
    if any(not source.authors or not source.year for source in sources if source.id in cited):
        full_actions.append("Complete missing bibliographic metadata before submission")
    if not cited:
        full_actions.append(
            "No sources are cited in the manuscript; strengthen evidence-grounded synthesis"
        )
    now_year = datetime.now(UTC).year
    used = [source for source in sources if source.id in cited]
    input_items = (store.stage("intake") or {}).get("output", {}).get("inputs", [])
    report = {
        "section_word_counts": {
            s.title: sum(len(p.text.split()) for p in s.paragraphs) for s in sections
        },
        "tables_included": ["literature-matrix.csv"]
        + (["computed-results.csv"] if results else []),
        "tables_delivery": "Separate evidence tables; author must integrate and number journal tables",
        "figures_proposed": (plan or {}).get("figures", []),
        "equation_labels_detected": re.findall(
            r"(?:\\tag\{(\d+)\}|\bEquation\s+(\d+))", manuscript
        ),
        "equation_check": "Detected labels only; symbol definitions, numbering and mathematical correctness require review",
        "datasets_used": [
            {"id": item["id"], "path": item["path"], "status": item["status"]}
            for item in input_items
            if item["role"] == "dataset"
        ],
        "statistical_analyses_performed": [
            {
                "result_id": result["id"],
                "calculator": result["calculator"],
                "assumptions": result.get("assumptions", []),
            }
            for result in results
            if result["calculator"] != "literature_inventory_v1"
        ],
        "unresolved_data_gaps": (plan or {}).get("missing_evidence", []),
        "claims_requiring_experimental_verification": [
            {"section": s.title, "text": p.text}
            for s in sections
            for p in s.paragraphs
            if p.kind == "proposal"
        ],
        "reference_count": len(cited),
        "recent_reference_percentage": 100
        * sum(bool(source.year and source.year >= now_year - recent_years + 1) for source in used)
        / len(used)
        if used
        else 0,
        "recent_year_window": [now_year - recent_years + 1, now_year],
        "references_requiring_metadata_verification": [
            s.id for s in used if not s.metadata_verified
        ],
        "uncited_discovered_sources": [s.id for s in sources if s.id not in cited],
        "abstract_only_sources": [s.id for s in used if s.access_level == "abstract"],
        "computed_results": results,
        "author_actions": full_actions,
        "review": review,
        "provider_usage": store.snapshot(),
        "submission_ready": False,
        "author_verification_required": True,
        "integrity_note": "Draft independently generated with traceable evidence; originality, claim entailment, source relevance and study authenticity still require author verification. No plagiarism/AI-detection guarantee.",
    }
    files.append(
        store.write(
            "outputs/quality-control.json", json.dumps(report, indent=2, ensure_ascii=False)
        )
    )
    files.append(
        store.write(
            "outputs/source-register.json",
            json.dumps([s.model_dump(mode="json") for s in sources], indent=2, ensure_ascii=False),
        )
    )
    files.append(
        store.write(
            "outputs/claim-register.json",
            json.dumps(
                [
                    {"section": s.title, **p.model_dump(mode="json")}
                    for s in sections
                    for p in s.paragraphs
                ],
                indent=2,
                ensure_ascii=False,
            ),
        )
    )
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    assessed = {item["source_id"]: item for item in (assessments or [])}
    writer.writerow(
        [
            "Source ID",
            "Year",
            "Title",
            "Access",
            "Method",
            "Dataset/System",
            "Finding",
            "Reported limitation",
            "Missing accessible detail",
        ]
    )
    for source in sources:
        appraisal = assessed.get(source.id, {})
        summaries = [
            (appraisal.get(field) or {}).get("summary", "[NOT REPORTED IN ACCESSIBLE TEXT]")
            for field in ("method", "dataset_or_system", "finding", "reported_limitation")
        ]
        writer.writerow(
            [
                source.id,
                source.year,
                source.title,
                source.access_level,
                *summaries,
                "; ".join(appraisal.get("missing_from_accessible_text", [])),
            ]
        )
    files.append(store.write("outputs/literature-matrix.csv", buffer.getvalue()))
    specifications = (plan or {}).get("figures", [])
    for number, spec in enumerate(specifications, 1):
        filename = f"figure-{number:02d}.svg"
        svg = svg_diagram(spec["nodes"], spec["edges"], f"Figure {number}: proposed architecture")
        files.append(store.write("outputs/" + filename, svg))
        full_actions.append(
            f"Verify proposed Figure {number}: {spec['title']} against the implemented study"
        )
    if specifications:
        files.append(
            store.write("outputs/figure-specifications.json", json.dumps(specifications, indent=2))
        )
    if used and report["recent_reference_percentage"] < 70:
        full_actions.append(
            "Recent citations are below the prompt's approximate 70–80% target; justify seminal sources or strengthen recent evidence"
        )
    # Add only genuine computed observations to result tables; no generated numerical values.
    if results:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["Result ID", "Input evidence ID", "Metric", "Value", "Unit"])
        for result in results:
            for metric, value in result["metrics"].items():
                writer.writerow(
                    [
                        result["id"],
                        result["input_id"],
                        metric,
                        "" if value is None else value,
                        "count" if metric.endswith("count") else "fraction",
                    ]
                )
        files.append(store.write("outputs/computed-results.csv", buffer.getvalue()))
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        full_actions.append("Install the documents extra to export DOCX")
    else:
        doc = Document()
        doc.styles["Normal"].font.name = "Times New Roman"
        doc.styles["Normal"].font.size = Pt(12)
        for line in manuscript.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:], 0)
            elif line.startswith("## "):
                doc.add_heading(line[3:], 1)
            elif line.strip():
                doc.add_paragraph(line)
        path = store.safe_path("outputs/manuscript.docx")
        temp = path.with_name("manuscript.tmp.docx")
        doc.save(temp)
        temp.replace(path)
        files.append(path)
    # Final report includes figure/export actions discovered during rendering.
    report["author_actions"] = list(dict.fromkeys(full_actions))
    store.write("outputs/quality-control.json", json.dumps(report, indent=2, ensure_ascii=False))
    return {
        "files": [path.relative_to(store.root).as_posix() for path in files],
        "author_actions": report["author_actions"],
        "submission_ready": False,
    }
