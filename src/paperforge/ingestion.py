from __future__ import annotations

import csv
import hashlib
import io
import json
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

import yaml

from paperforge.author_validation import load_author_validation, validation_evidence
from paperforge.config import IngestionConfig
from paperforge.domain import AuthorValidationDecision, EvidenceItem, EvidenceKind
from paperforge.storage import ProjectStore

TEXT_EXTENSIONS = {".txt", ".md", ".rst", ".tex", ".bib", ".ris", ".yaml", ".yml"}
FIGURE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".svg"}
RESPONSE_FILENAMES = {"responses.yaml", "responses.yml", "respones.yaml", "respones.yml"}


@dataclass(slots=True)
class IngestionReport:
    discovered: int = 0
    extracted: int = 0
    unchanged: int = 0
    skipped: int = 0
    warnings: list[str] = field(default_factory=list)
    fingerprint: str = ""


class DocumentIngestor:
    """Extract user-controlled project files into a provenance-preserving evidence registry."""

    def __init__(self, store: ProjectStore, config: IngestionConfig) -> None:
        self.store = store
        self.config = config

    def refresh(self) -> IngestionReport:
        report = IngestionReport()
        paths = self.store.input_files()
        report.discovered = len(paths)
        existing = self.store.load_evidence()
        generated = [item for item in existing if item.metadata.get("generated")]
        previous = {
            item.source_path: item
            for item in existing
            if item.metadata.get("auto_ingested") and item.source_path
        }
        ingested: list[EvidenceItem] = []
        fingerprint = hashlib.sha256()
        remaining = self.config.max_total_chars

        for path in paths:
            relative = path.relative_to(self.store.root).as_posix()
            size = path.stat().st_size
            if size > self.config.max_file_bytes:
                report.skipped += 1
                report.warnings.append(f"Skipped oversized input: {relative}")
                continue
            checksum = self.store.checksum(path)
            fingerprint.update(relative.encode("utf-8"))
            fingerprint.update(checksum.encode("ascii"))
            old = previous.get(relative)
            if old and old.checksum == checksum:
                ingested.append(old)
                remaining -= len(old.content)
                report.unchanged += 1
                continue
            try:
                content, locator = self._extract(path)
            except (OSError, ValueError, RuntimeError) as exc:
                report.skipped += 1
                report.warnings.append(f"Could not extract {relative}: {exc}")
                continue
            if not content.strip():
                report.skipped += 1
                report.warnings.append(f"Skipped empty input: {relative}")
                continue
            limit = max(0, min(self.config.max_chars_per_document, remaining))
            if limit == 0:
                report.skipped += 1
                report.warnings.append(
                    "Evidence context limit reached; remaining files were indexed only."
                )
                continue
            truncated = len(content) > limit
            content = content[:limit]
            remaining -= len(content)
            ingested.append(
                EvidenceItem(
                    id=self._evidence_id(relative),
                    kind=self._kind_for(relative, path.suffix.casefold()),
                    title=path.name,
                    content=content,
                    source_path=relative,
                    locator=locator,
                    checksum=checksum,
                    verified=True,
                    metadata={
                        "auto_ingested": True,
                        "verification": "user_supplied_file_with_sha256",
                        "size_bytes": size,
                        "extension": path.suffix.casefold(),
                        "truncated": truncated,
                    },
                )
            )
            report.extracted += 1

        if self.store.author_validation_path.exists():
            package = load_author_validation(self.store.author_validation_path)
            report.discovered += 1
            validation_item = validation_evidence(package)
            if validation_item:
                ingested.append(validation_item)
                fingerprint.update(validation_item.checksum.encode("ascii"))
                report.extracted += 1
            for item in package.items:
                if item.decision == AuthorValidationDecision.PENDING and item.answer:
                    report.warnings.append(
                        f"Ignored {item.id} answer because its decision is still 'pending'."
                    )
                if (
                    item.decision
                    in {
                        AuthorValidationDecision.PROVIDED,
                        AuthorValidationDecision.CORRECTED,
                    }
                    and not item.answer
                ):
                    report.warnings.append(
                        f"Ignored {item.id} because decision '{item.decision.value}' requires an answer."
                    )

        self.store.save_evidence(generated + ingested)
        report.fingerprint = fingerprint.hexdigest()
        return report

    def _extract(self, path: Path) -> tuple[str, str]:
        if path.name.casefold() in RESPONSE_FILENAMES:
            return self._extract_responses(path)
        suffix = path.suffix.casefold()
        extractors: dict[str, Callable[[Path], tuple[str, str]]] = {
            ".json": self._extract_json,
            ".csv": self._extract_delimited,
            ".tsv": self._extract_delimited,
            ".pdf": self._extract_pdf,
            ".docx": self._extract_docx,
            ".xlsx": self._extract_xlsx,
        }
        if suffix in TEXT_EXTENSIONS:
            return path.read_text(encoding="utf-8", errors="replace"), "full file"
        if suffix in extractors:
            return extractors[suffix](path)
        if suffix in FIGURE_EXTENSIONS:
            return (
                f"Figure file supplied: {path.name}. The image is registered as an artifact, "
                "but its pixels are not interpreted as scientific evidence without a caption.",
                "file metadata",
            )
        raise ValueError(f"unsupported file type '{suffix or '[none]'}'")

    @staticmethod
    def _extract_responses(path: Path) -> tuple[str, str]:
        try:
            payload = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        except yaml.YAMLError as exc:
            raise ValueError(f"invalid responses YAML: {exc}") from exc
        answers = payload.get("answers", payload) if isinstance(payload, dict) else {}
        if not isinstance(answers, dict):
            raise ValueError("responses.yaml must contain an 'answers' mapping")
        supplied = [
            f"{identifier}: {str(answer).strip()}"
            for identifier, answer in answers.items()
            if answer is not None and str(answer).strip()
        ]
        if not supplied:
            raise ValueError("responses.yaml contains no non-empty author responses")
        return (
            "Author-supplied factual responses from the legacy intake workflow:\n\n"
            + "\n\n".join(supplied),
            "all non-empty answers",
        )

    @staticmethod
    def _extract_json(path: Path) -> tuple[str, str]:
        value = json.loads(path.read_text(encoding="utf-8"))
        return json.dumps(value, indent=2, ensure_ascii=False), "JSON document"

    @staticmethod
    def _extract_delimited(path: Path) -> tuple[str, str]:
        delimiter = "\t" if path.suffix.casefold() == ".tsv" else ","
        output = io.StringIO()
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as source:
            reader = csv.reader(source, delimiter=delimiter)
            writer = csv.writer(output, delimiter=delimiter, lineterminator="\n")
            for row_number, row in enumerate(reader, start=1):
                writer.writerow(row)
                if row_number >= 20_000:
                    break
        return output.getvalue(), "rows 1-20000"

    @staticmethod
    def _extract_pdf(path: Path) -> tuple[str, str]:
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("install PaperForge with [documents] for PDF extraction") from exc
        chunks: list[str] = []
        with fitz.open(path) as document:
            page_count = len(document)
            for index, page in enumerate(document, start=1):
                chunks.append(f"\n--- Page {index} ---\n{page.get_text('text')}")
        return "".join(chunks), f"pages 1-{page_count}"

    @staticmethod
    def _extract_docx(path: Path) -> tuple[str, str]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("install PaperForge with [documents] for DOCX extraction") from exc
        document = Document(path)
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_number, table in enumerate(document.tables, start=1):
            chunks.append(f"\n[Table {table_number}]")
            chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(chunks), "paragraphs and tables"

    @staticmethod
    def _extract_xlsx(path: Path) -> tuple[str, str]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("install PaperForge with [documents] for XLSX extraction") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        output = io.StringIO()
        writer = csv.writer(output, lineterminator="\n")
        for sheet in workbook.worksheets:
            writer.writerow([f"[Sheet: {sheet.title}]"])
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                writer.writerow(["" if value is None else value for value in row])
                if row_number >= 20_000:
                    break
        workbook.close()
        return output.getvalue(), "all sheets, rows 1-20000 per sheet"

    @staticmethod
    def _evidence_id(relative_path: str) -> str:
        digest = hashlib.sha256(relative_path.casefold().encode("utf-8")).hexdigest()[:12].upper()
        return f"EV-{digest}"

    @staticmethod
    def _kind_for(relative_path: str, suffix: str) -> EvidenceKind:
        if relative_path.startswith("data/") or suffix in {".csv", ".tsv", ".xlsx"}:
            return EvidenceKind.EXPERIMENTAL_DATA
        if relative_path.startswith("figures/") or suffix in FIGURE_EXTENSIONS:
            return EvidenceKind.FIGURE
        if relative_path.startswith("inputs/"):
            return EvidenceKind.USER_FACT
        return EvidenceKind.SOURCE_DOCUMENT
