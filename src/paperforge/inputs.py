from __future__ import annotations

import csv
import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

from paperforge.config import Ingestion
from paperforge.schemas import InputItem
from paperforge.store import Store, fingerprint

TEXT = {".txt", ".md", ".rst", ".tex", ".bib", ".ris", ".yaml", ".yml", ".json", ".m", ".py", ".r"}
IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}
NATIVE = {".cdr", ".slx", ".mdl", ".mlx", ".fig", ".dwg", ".xls", ".ppt"}
ROLES = {"author_note", "dataset", "source", "figure", "code", "artifact"}


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def input_signature(store: Store) -> str:
    records = []
    for path in sorted((store.root / "inputs").rglob("*")):
        if path.is_symlink():
            raise ValueError("Input symlinks are not supported; attach a copy inside the project")
        if path.is_file():
            records.append((path.relative_to(store.root).as_posix(), checksum(path)))
    return fingerprint(records)


def archive_guard(path: Path, limit: int):
    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > limit:
            raise ValueError("Expanded document exceeds configured archive limit")


class Ingestor:
    def __init__(self, store: Store, settings: Ingestion):
        self.store, self.settings = store, settings

    def run(self) -> list[InputItem]:
        manifest_path = self.store.root / "inputs" / "manifest.json"
        manifest = (
            json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
        )
        if not isinstance(manifest, dict) or any(
            not isinstance(value, str) or value not in ROLES for value in manifest.values()
        ):
            raise ValueError(
                "inputs/manifest.json must map input-relative paths to valid evidence roles"
            )
        items = []
        for path in sorted((self.store.root / "inputs").rglob("*")):
            if path.is_symlink() or not path.resolve().is_relative_to(self.store.root):
                raise ValueError("Input symlinks or paths outside the project are not supported")
            if (
                not path.is_file()
                or path == manifest_path
                or path == self.store.root / "inputs" / "source-dois.json"
            ):
                continue
            relative = path.relative_to(self.store.root).as_posix()
            suffix = path.suffix.lower()
            role = manifest.get(
                path.relative_to(self.store.root / "inputs").as_posix(), self._role(suffix)
            )
            digest = checksum(path)
            cache_key = "input:" + fingerprint(
                [relative, digest, role, self.settings.model_dump(mode="json"), 1]
            )
            if cached := self.store.cached_checkpoint(cache_key):
                items.append(InputItem.model_validate(cached))
                continue
            item = InputItem(
                id="EV-" + hashlib.sha256(relative.encode()).hexdigest()[:12],
                path=relative,
                sha256=digest,
                role=role,
                status="extracted",
                metadata={
                    "size_bytes": path.stat().st_size,
                    "verification": "file provenance only; scientific validity is not established",
                },
            )
            if path.stat().st_size > self.settings.max_file_bytes:
                item.status = "unsupported"
                item.warnings.append(
                    "File exceeds max_file_bytes; increase limit or supply a smaller export"
                )
            else:
                try:
                    item.chunks, item.warnings, item.metadata = self.extract(path, item.metadata)
                    if suffix in NATIVE:
                        item.status = "conversion_required"
                    elif not item.chunks:
                        item.status = (
                            "partial" if suffix in IMAGES | {".pdf", ".svg"} else "unsupported"
                        )
                    elif item.warnings:
                        item.status = "partial"
                    remaining = self.settings.max_extracted_chars
                    limited = []
                    for chunk in item.chunks:
                        if remaining <= 0:
                            break
                        limited.append(
                            {"locator": chunk["locator"], "text": chunk["text"][:remaining]}
                        )
                        remaining -= len(limited[-1]["text"])
                    if sum(len(c["text"]) for c in item.chunks) > self.settings.max_extracted_chars:
                        item.status = "partial"
                        item.warnings.append(
                            "Extraction truncated at max_extracted_chars; original retained"
                        )
                    item.chunks = limited
                except Exception as exc:
                    # Keep one bad attachment from discarding successfully read independent inputs.
                    item.status = "failed"
                    item.warnings.append(
                        f"{type(exc).__name__}: extraction failed; install the relevant extra or supply a supported export"
                    )
            self.store.checkpoint(cache_key, item.model_dump(mode="json"))
            items.append(item)
        return items

    @staticmethod
    def _role(suffix: str) -> str:
        if suffix in {".csv", ".tsv", ".xlsx", ".mat", ".npy"}:
            return "dataset"
        if suffix in IMAGES | {".svg", ".cdr", ".fig"}:
            return "figure"
        if suffix in {".m", ".py", ".r", ".slx", ".mdl", ".mlx"}:
            return "code"
        if suffix in {".pdf", ".docx", ".pptx"}:
            return "source"
        return "author_note"

    def extract(self, path: Path, metadata: dict) -> tuple[list[dict], list[str], dict]:
        suffix = path.suffix.lower()
        warnings = []
        chunks = []
        if suffix in TEXT:
            text = path.read_text(encoding="utf-8-sig")
            chunks = [{"locator": "full file", "text": text}]
            if suffix in {".m", ".py", ".r"}:
                warnings.append("Code is read as text; it is never executed during ingestion")
        elif suffix in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as file:
                for row_id, row in enumerate(
                    csv.reader(file, delimiter="\t" if suffix == ".tsv" else ","), 1
                ):
                    if row_id > self.settings.max_rows:
                        warnings.append("Row limit reached")
                        break
                    chunks.append(
                        {"locator": f"row {row_id}", "text": json.dumps(row, ensure_ascii=False)}
                    )
        elif suffix == ".xlsx":
            archive_guard(path, self.settings.max_archive_bytes)
            from openpyxl import load_workbook

            formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            cells = 0
            try:
                for sheet in formulas:
                    rows = zip(sheet.iter_rows(), values[sheet.title].iter_rows(), strict=True)
                    for row_id, (row, cached_row) in enumerate(rows, 1):
                        if row_id > self.settings.max_rows or cells >= self.settings.max_cells:
                            warnings.append("Workbook row/cell limit reached")
                            break
                        for cell, cached_cell in zip(row, cached_row, strict=True):
                            cells += 1
                            if cells > self.settings.max_cells:
                                break
                            if cell.value is None:
                                continue
                            value = cached_cell.value if cell.data_type == "f" else cell.value
                            text = f"value={value!s}"
                            if cell.data_type == "f":
                                text += f"; formula={cell.value}; cached value may be missing or stale (not recalculated)"
                            chunks.append(
                                {"locator": f"{sheet.title}!{cell.coordinate}", "text": text}
                            )
                warnings.append(
                    "Excel formulas are not recalculated; macros and external links are not executed"
                )
            finally:
                formulas.close()
                values.close()
        elif suffix == ".pdf":
            import fitz

            with fitz.open(path) as document:
                for number, page in enumerate(document, 1):
                    if number > 200:
                        warnings.append("PDF page limit reached")
                        break
                    text = page.get_text()
                    if text.strip():
                        chunks.append({"locator": f"page {number}", "text": text})
                    else:
                        warnings.append(
                            f"Page {number} has no extractable text; supply OCR text or a caption"
                        )
        elif suffix == ".docx":
            archive_guard(path, self.settings.max_archive_bytes)
            from docx import Document

            doc = Document(path)
            chunks = [
                {"locator": f"paragraph {i}", "text": p.text}
                for i, p in enumerate(doc.paragraphs, 1)
                if p.text.strip()
            ]
            for i, table in enumerate(doc.tables, 1):
                for j, row in enumerate(table.rows, 1):
                    chunks.append(
                        {
                            "locator": f"table {i} row {j}",
                            "text": " | ".join(cell.text for cell in row.cells),
                        }
                    )
        elif suffix == ".pptx":
            archive_guard(path, self.settings.max_archive_bytes)
            from pptx import Presentation

            presentation = Presentation(path)
            for i, slide in enumerate(presentation.slides, 1):
                for j, shape in enumerate(slide.shapes, 1):
                    if shape.has_text_frame:
                        chunks.append({"locator": f"slide {i} shape {j}", "text": shape.text})
                    if shape.has_table:
                        for k, row in enumerate(shape.table.rows, 1):
                            chunks.append(
                                {
                                    "locator": f"slide {i} table {j} row {k}",
                                    "text": " | ".join(c.text for c in row.cells),
                                }
                            )
            warnings.append(
                "Slide pictures and charts require source data/captions for scientific interpretation"
            )
        elif suffix in IMAGES:
            from PIL import Image

            with Image.open(path) as image:
                metadata.update(
                    {"width": image.width, "height": image.height, "format": image.format}
                )
                image.verify()
            if self.settings.ocr and shutil.which("tesseract"):
                result = subprocess.run(
                    ["tesseract", str(path), "stdout"],
                    capture_output=True,
                    text=True,
                    timeout=30,
                    check=True,
                )
                if result.stdout.strip():
                    chunks = [{"locator": "OCR text (unverified)", "text": result.stdout}]
            warnings.append(
                "Image registered; OCR/visual appearance is not verified measurement data. Supply caption and raw values"
            )
        elif suffix == ".svg":
            from defusedxml import ElementTree

            root = ElementTree.fromstring(path.read_bytes())
            chunks = [
                {
                    "locator": "SVG labels",
                    "text": " ".join(
                        "".join(node.itertext())
                        for node in root.iter()
                        if node.tag.endswith("}text")
                    ),
                }
            ]
            warnings.append(
                "SVG labels extracted; geometry and numerical interpretation require author verification"
            )
        elif suffix in {".mat", ".npy"}:
            import numpy as np

            if suffix == ".npy":
                arrays = {"array": np.load(path, allow_pickle=False, mmap_mode="r")}
            else:
                import h5py
                from scipy.io import loadmat

                if h5py.is_hdf5(path):
                    with h5py.File(path) as file:
                        for name, value in file.items():
                            if (
                                isinstance(value, h5py.Dataset)
                                and value.dtype.kind in "biufc"
                                and value.size <= self.settings.max_cells
                            ):
                                chunks.append(
                                    {
                                        "locator": f"variable {name}",
                                        "text": f"shape={value.shape}; values={np.array2string(value[()], threshold=self.settings.max_cells, precision=17, floatmode='unique')}",
                                    }
                                )
                            else:
                                warnings.append(
                                    f"Variable {name} needs specialized MATLAB object/reference handling or exceeds cell limit"
                                )
                    arrays = {}
                else:
                    arrays = loadmat(path)
            for name, value in arrays.items():
                if name.startswith("__"):
                    continue
                if (
                    isinstance(value, np.ndarray)
                    and value.dtype.kind in "biufc"
                    and value.size <= self.settings.max_cells
                ):
                    chunks.append(
                        {
                            "locator": f"variable {name}",
                            "text": f"shape={value.shape}; values={np.array2string(value, threshold=self.settings.max_cells, precision=17, floatmode='unique')}",
                        }
                    )
                else:
                    warnings.append(
                        f"Variable {name} needs a simpler numeric export or exceeds cell limit"
                    )
        elif suffix in NATIVE:
            warnings.append(
                "Native file retained. Export CDR/FIG as SVG or PDF; export simulation data as CSV/MAT and model description as text/PDF. Execution requires a separately configured compatible runtime"
            )
            sibling_exports = [
                p.name
                for p in path.parent.glob(path.stem + ".*")
                if p.suffix.lower() in {".svg", ".pdf", ".png", ".csv"}
            ]
            metadata["available_exports"] = sibling_exports
        else:
            warnings.append(
                f"No adapter for {suffix or 'extensionless file'}; supply a supported export"
            )
        return chunks, warnings, metadata
