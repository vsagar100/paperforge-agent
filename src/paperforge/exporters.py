from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from paperforge.author_validation import author_validation_markdown
from paperforge.citations import build_bibtex, render_numbered_citations
from paperforge.domain import EvidenceCoverageStatus, IssueDisposition, WorkflowState
from paperforge.standards import normalize_heading
from paperforge.storage import ProjectStore


@dataclass(slots=True)
class ExportReport:
    files: list[Path] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class OutputExporter:
    def __init__(self, store: ProjectStore) -> None:
        self.store = store

    def export(self, state: WorkflowState) -> ExportReport:
        report = ExportReport()
        references = self.store.load_references()
        rendered, cited = render_numbered_citations(self.store.read_manuscript(), references)
        try:
            profile = self.store.load_publication_profile()
        except (FileNotFoundError, ValueError):
            profile = None
        if profile and profile.number_sections:
            rendered = self._number_section_headings(rendered)
        self.store.write_text("outputs/manuscript.md", rendered)
        report.files.append(self.store.root / "outputs" / "manuscript.md")

        self.store.write_text("outputs/references.bib", build_bibtex(cited))
        report.files.append(self.store.root / "outputs" / "references.bib")

        self.store.write_text("outputs/literature-matrix.csv", self._literature_matrix())
        report.files.append(self.store.root / "outputs" / "literature-matrix.csv")

        quality = self._quality_payload(state, rendered, cited)
        self.store.write_json("outputs/quality-report.json", quality)
        self.store.write_text("outputs/review-report.md", self._review_report(quality))
        report.files.extend(
            [
                self.store.root / "outputs" / "quality-report.json",
                self.store.root / "outputs" / "review-report.md",
            ]
        )

        self.store.write_text("outputs/revision-history.md", self._revision_history(state))
        report.files.append(self.store.root / "outputs" / "revision-history.md")

        self.store.write_text("outputs/publication-profile.md", self._publication_profile())
        self.store.write_text("outputs/evidence-coverage.md", self._evidence_coverage())
        self.store.write_text("outputs/display-item-plan.md", self._display_item_plan())
        self.store.write_text("outputs/submission-checklist.md", self._submission_checklist(state))
        report.files.extend(
            [
                self.store.root / "outputs" / "publication-profile.md",
                self.store.root / "outputs" / "evidence-coverage.md",
                self.store.root / "outputs" / "display-item-plan.md",
                self.store.root / "outputs" / "submission-checklist.md",
            ]
        )

        if self.store.author_validation_path.exists():
            self.store.write_text(
                "outputs/author-validation.md",
                author_validation_markdown(self.store.load_author_validation()),
            )
            report.files.append(self.store.root / "outputs" / "author-validation.md")

        try:
            docx = self._export_docx(rendered, profile)
        except ImportError:
            report.warnings.append(
                "DOCX export skipped; install PaperForge with the [documents] extra."
            )
        else:
            report.files.append(docx)
        return report

    def _publication_profile(self) -> str:
        try:
            profile = self.store.load_publication_profile()
        except (FileNotFoundError, ValueError):
            return "# Publication Profile\n\nNo publication profile is available.\n"
        lines = [
            "# Publication Profile",
            "",
            f"- Profile: `{profile.profile_id}`",
            f"- Journal: {profile.journal_name or 'Target journal not specified'}",
            f"- Source: {profile.source_label}",
            f"- Target-journal rules verified: {'yes' if profile.target_rules_verified else 'no'}",
            f"- Article type: `{profile.article_type}`",
            f"- Main-text range: {profile.word_min}–{profile.word_max or 'not specified'} words",
            f"- Abstract: {profile.abstract_min_words}–{profile.abstract_max_words} words",
            f"- Keywords: {profile.keyword_min}–{profile.keyword_max}",
            f"- References: {profile.reference_min}–{profile.reference_max or 'not specified'}",
            f"- Tables: maximum {profile.table_max if profile.table_max is not None else 'not specified'}",
            f"- Figures: maximum {profile.figure_max if profile.figure_max is not None else 'not specified'}",
            "",
            "## Indexing context",
            "",
            profile.indexing_context,
            "",
            "## Rules applied",
            "",
            *[f"- {item}" for item in profile.formatting_rules],
            "",
            "## Sources",
            "",
            *[f"- {item}" for item in profile.source_urls],
            "",
        ]
        return "\n".join(lines)

    def _evidence_coverage(self) -> str:
        try:
            coverage = self.store.load_evidence_coverage()
        except (FileNotFoundError, ValueError):
            return "# Evidence Coverage\n\nNo evidence coverage report is available.\n"
        lines = ["# Evidence Coverage", ""]
        for item in coverage.requirements:
            lines.extend(
                [
                    f"## {item.label}",
                    "",
                    f"- Code: `{item.code}`",
                    f"- Level: `{item.level.value}`",
                    f"- Status: `{item.status.value}`",
                    f"- Evidence atoms: {', '.join(item.matched_claim_ids) or 'none located'}",
                    f"- Assessment: {item.explanation}",
                    f"- Required detail: {item.requested_detail or 'none'}",
                    "",
                ]
            )
        return "\n".join(lines)

    def _display_item_plan(self) -> str:
        try:
            outline = self.store.load_outline()
        except (FileNotFoundError, ValueError):
            return "# Display Item Plan\n\nNo display-item plan is available.\n"
        lines = [
            "# Display Item Plan",
            "",
            "Only items marked `available` may be generated from registered evidence. "
            "Items marked `author_required` require an authentic source file or author-approved design.",
            "",
        ]
        if not outline.display_items:
            lines.append("No table or figure is currently planned.")
            return "\n".join(lines) + "\n"
        for item in outline.display_items:
            lines.extend(
                [
                    f"## {item.id}: {item.title}",
                    "",
                    f"- Type: `{item.kind}`",
                    f"- Section: {item.section}",
                    f"- Status: `{item.status}`",
                    f"- Purpose: {item.purpose}",
                    f"- Claim atoms: {', '.join(item.claim_ids) or 'none'}",
                    f"- Evidence records: {', '.join(item.evidence_ids) or 'none'}",
                    "",
                ]
            )
        return "\n".join(lines)

    def _submission_checklist(self, state: WorkflowState) -> str:
        try:
            profile = self.store.load_publication_profile()
        except (FileNotFoundError, ValueError):
            profile = None
        unresolved = [action for action in state.author_actions if not action.resolved]
        pending_validation = []
        if self.store.author_validation_path.exists():
            pending_validation = self.store.load_author_validation().pending_items
        lines = [
            "# Submission Checklist",
            "",
            f"- Workflow complete: {'yes' if state.workflow_completed else 'no'}",
            f"- Submission candidate: {'yes' if state.submission_ready else 'no'}",
            "- Scientific content and all declarations reviewed by every author: pending author confirmation",
            "- Title page, affiliations, ORCIDs, and corresponding-author details: pending author confirmation",
            "- Main manuscript anonymized for the journal review model: pending author confirmation",
            "- Figures checked against original files and resolution requirements: pending author confirmation",
            "- Tables, equations, units, and cross-references checked: pending author confirmation",
            "- Reference metadata and DOI links checked against source records: pending author confirmation",
            "- Similarity/originality check performed by an authorized service: not performed by PaperForge",
            f"- Consolidated study-fact validation: {len(pending_validation)} pending item(s)",
            "",
            "## Unresolved author actions",
            "",
        ]
        lines.extend(
            [f"- {item.action} — {item.reason}" for item in unresolved]
            or ["- None recorded by the implemented gates."]
        )
        if profile:
            lines.extend(
                [
                    "",
                    "## Target-journal submission files",
                    "",
                    "- Editable anonymized Word manuscript",
                    "- Separate title page",
                    "- Cover letter",
                    "- Author-approved declarations and permissions",
                    "- Original figure files and any supplementary/repository materials",
                    "",
                    f"Recheck the live journal instructions before submission: {profile.source_urls[0] if profile.source_urls else 'not configured'}",
                ]
            )
        return "\n".join(lines) + "\n"

    def _literature_matrix(self) -> str:
        references = self.store.load_references()
        try:
            synthesis = self.store.load_synthesis()
            notes = {note.reference_id: note for note in synthesis.source_notes}
        except (FileNotFoundError, ValueError):
            notes = {}
        output = io.StringIO()
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(
            [
                "Reference ID",
                "Authors",
                "Year",
                "Title",
                "Venue",
                "Volume",
                "Issue",
                "Pages",
                "DOI",
                "Method",
                "Dataset or material",
                "Key finding",
                "Limitation",
                "Relevance",
                "Verification",
            ]
        )
        for reference in references:
            note = notes.get(reference.id)
            writer.writerow(
                [
                    reference.id,
                    "; ".join(reference.authors),
                    reference.year or "",
                    reference.title,
                    reference.venue or "",
                    reference.volume or "",
                    reference.issue or "",
                    reference.pages or "",
                    reference.doi or "",
                    note.method if note else "Not reported in the available metadata.",
                    note.dataset_or_material if note else "Not reported in the available metadata.",
                    note.key_finding if note else "Not reported in the available metadata.",
                    note.limitation if note else "Not reported in the available metadata.",
                    note.relevance if note else "",
                    "; ".join(reference.verification_sources),
                ]
            )
        return output.getvalue()

    def _quality_payload(
        self,
        state: WorkflowState,
        rendered_manuscript: str,
        cited,
    ) -> dict:
        records = {
            stage: record.model_dump(mode="json") for stage, record in state.stage_records.items()
        }
        unresolved = [
            issue.model_dump(mode="json")
            for record in state.stage_records.values()
            for issue in record.issues
            if not issue.resolved
        ]
        blockers = [
            issue
            for issue in unresolved
            if issue.get("disposition") == IssueDisposition.INTEGRITY_BLOCKER
        ]
        return {
            "schema_version": 3,
            "project_id": state.project_id,
            "paper_type": state.profile.resolved_paper_type,
            "target_journal": state.profile.target_journal,
            "workflow_completed": state.workflow_completed,
            "submission_ready": state.submission_ready,
            "readiness_label": (
                "submission_candidate" if state.submission_ready else "author_action_required"
            ),
            "word_count": len(
                re.findall(
                    r"\b[\w'-]+\b",
                    re.split(r"(?im)^##\s+(?:\d+\.\s+)?References\s*$", rendered_manuscript)[0],
                )
            ),
            "cited_sources": len(cited),
            "verified_cited_sources": sum(reference.verified for reference in cited),
            "section_word_counts": self._section_word_counts(rendered_manuscript),
            "table_count": len(
                re.findall(
                    r"(?m)^\|(?:[^\n|]+\|)+\s*$\n^\|\s*:?-{3,}",
                    rendered_manuscript,
                )
            ),
            "evidence_coverage": self._coverage_payload(),
            "integrity_blockers": blockers,
            "unresolved_review_items": unresolved,
            "author_actions": [action.model_dump(mode="json") for action in state.author_actions],
            "author_validation": self._author_validation_payload(),
            "stage_records": records,
            "disclaimer": (
                "Submission-ready means the implemented evidence and consistency gates passed. "
                "It does not guarantee journal acceptance and does not replace author verification."
            ),
        }

    def _author_validation_payload(self) -> dict:
        if not self.store.author_validation_path.exists():
            return {"available": False, "items": 0, "pending": 0}
        package = self.store.load_author_validation()
        return {
            "available": True,
            "items": len(package.items),
            "pending": len(package.pending_items),
            "path": "author-actions/validation.yaml",
        }

    @staticmethod
    def _review_report(quality: dict) -> str:
        lines = [
            "# PaperForge Review Report",
            "",
            f"- Readiness: **{quality['readiness_label']}**",
            f"- Manuscript words: {quality['word_count']}",
            f"- Cited sources: {quality['cited_sources']}",
            f"- Verified cited sources: {quality['verified_cited_sources']}",
            "",
            "## Integrity blockers",
            "",
        ]
        blockers = quality["integrity_blockers"]
        if blockers:
            lines.extend(
                f"- **{issue['code']}** ({issue.get('section') or 'Manuscript'}): "
                f"{issue['description']}"
                for issue in blockers
            )
        else:
            lines.append("- None detected by the implemented gates.")
        lines.extend(["", "## Author actions and recommendations", ""])
        items = [issue for issue in quality["unresolved_review_items"] if issue not in blockers]
        if items:
            lines.extend(
                f"- **{issue['code']}** ({issue.get('section') or 'Manuscript'}): "
                f"{issue['required_change']}"
                for issue in items
            )
        else:
            lines.append("- No unresolved review item.")
        lines.extend(["", "## Interpretation", "", quality["disclaimer"], ""])
        return "\n".join(lines)

    @staticmethod
    def _revision_history(state: WorkflowState) -> str:
        lines = ["# Revision History", ""]
        for record in state.run_history:
            lines.extend(
                [
                    f"## {record.stage} — {record.status.value}",
                    "",
                    f"- Completed: {record.completed_at or 'not recorded'}",
                    f"- Score: {record.score:.2f}",
                    f"- Model: {record.model or 'deterministic'}",
                ]
            )
            if record.changes:
                lines.append("- Changes:")
                lines.extend(f"  - {change}" for change in record.changes)
            if record.issues:
                lines.append(f"- Unresolved issues recorded: {len(record.issues)}")
            lines.append("")
        return "\n".join(lines)

    def _export_docx(self, manuscript: str, profile=None) -> Path:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Cm, Inches, Pt, RGBColor
        except ImportError as exc:
            raise ImportError from exc

        document = Document()
        section = document.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        styles = document.styles
        styles["Normal"].font.name = "Times New Roman"
        styles["Normal"].font.size = Pt(
            12 if profile and profile.profile_id.startswith("djes") else 11
        )
        styles["Normal"].paragraph_format.line_spacing = 1.15
        for level, size in {1: 16, 2: 12, 3: 11}.items():
            style = styles[f"Heading {level}"]
            style.font.name = "Times New Roman"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(9 if level > 1 else 12)
            style.paragraph_format.space_after = Pt(4)
        lines = manuscript.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if not stripped:
                index += 1
                continue
            if (
                stripped.startswith("|")
                and index + 1 < len(lines)
                and self._is_table_separator(lines[index + 1])
            ):
                headers = self._table_cells(stripped)
                rows: list[list[str]] = []
                index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append(self._table_cells(lines[index]))
                    index += 1
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                for column, value in enumerate(headers):
                    paragraph = table.rows[0].cells[column].paragraphs[0]
                    run = paragraph.add_run(value)
                    run.bold = True
                for values in rows:
                    cells = table.add_row().cells
                    for column, value in enumerate(values[: len(cells)]):
                        self._add_markdown_runs(cells[column].paragraphs[0], value)
                continue
            heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
            if heading:
                level = min(len(heading.group(1)), 3)
                paragraph = document.add_heading("", level=level)
                self._add_markdown_runs(paragraph, heading.group(2))
                if level == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                index += 1
                continue
            numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
            if numbered:
                paragraph = document.add_paragraph()
                self._add_markdown_runs(paragraph, stripped)
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
                index += 1
                continue
            bullet = re.match(r"^[-*]\s+(.+)$", stripped)
            if bullet:
                paragraph = document.add_paragraph(style="List Bullet")
                self._add_markdown_runs(paragraph, bullet.group(1))
                index += 1
                continue
            paragraph = document.add_paragraph(stripped)
            paragraph.clear()
            self._add_markdown_runs(paragraph, stripped)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.space_after = Pt(6)
            index += 1
        output = self.store.root / "outputs" / "manuscript.docx"
        document.save(output)
        return output

    def _coverage_payload(self) -> list[dict]:
        try:
            coverage = self.store.load_evidence_coverage()
        except (FileNotFoundError, ValueError):
            return []
        return [
            {
                "code": item.code,
                "level": item.level,
                "status": item.status,
                "supported": item.status == EvidenceCoverageStatus.SUPPORTED,
            }
            for item in coverage.requirements
        ]

    @staticmethod
    def _section_word_counts(manuscript: str) -> dict[str, int]:
        matches = list(re.finditer(r"(?m)^##\s+(.+?)\s*$", manuscript))
        counts: dict[str, int] = {}
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(manuscript)
            body = manuscript[match.end() : end]
            counts[match.group(1).strip()] = len(re.findall(r"\b[\w'-]+\b", body))
        return counts

    @staticmethod
    def _number_section_headings(manuscript: str) -> str:
        exempt = {
            "abstract",
            "keywords",
            "acknowledgements",
            "data availability",
            "conflict of interest",
            "funding",
            "author contributions",
            "declaration of ai use",
            "references",
        }
        main = 0
        sub = 0
        active_main: int | None = None
        output: list[str] = []
        for line in manuscript.splitlines():
            h2 = re.match(r"^##\s+(.+?)\s*$", line)
            if h2:
                heading = re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", h2.group(1)).strip()
                if normalize_heading(heading) in exempt:
                    active_main = None
                    output.append(f"## {heading}")
                else:
                    main += 1
                    sub = 0
                    active_main = main
                    output.append(f"## {main}. {heading}")
                continue
            h3 = re.match(r"^###\s+(.+?)\s*$", line)
            if h3 and active_main is not None:
                heading = re.sub(r"^\d+(?:\.\d+)*[.)]?\s*", "", h3.group(1)).strip()
                sub += 1
                output.append(f"### {active_main}.{sub}. {heading}")
                continue
            output.append(line)
        return "\n".join(output).rstrip() + "\n"

    @staticmethod
    def _add_markdown_runs(paragraph, value: str) -> None:
        for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", value):
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("*") and token.endswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            else:
                paragraph.add_run(token)

    @staticmethod
    def _table_cells(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip().strip("|").split("|")]

    @classmethod
    def _is_table_separator(cls, line: str) -> bool:
        cells = cls._table_cells(line)
        return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)
